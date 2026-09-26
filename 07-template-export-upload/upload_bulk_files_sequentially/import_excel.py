"""Read document records from the export Excel file."""

import logging
from pathlib import Path

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

NAME_COLUMN = "Name"  # grid header label on the target site, copied into the export
REQUIRED_COLUMNS = {NAME_COLUMN, "Document URL", "Document Path"}


def read_document_records(filename: str, sheet_name: str) -> list[dict]:
    """Read Excel and return list of document records for upload.

    Each record is a dict with keys:
        - name (str): Document name from the NAME_COLUMN column
        - url (str): Direct URL to document detail page from "Document URL"
        - file_path (str): Local file path from "Document Path"
        - row_number (int): Excel row number (for error reporting)

    Skips rows where URL or file_path is empty.
    """
    path = Path(filename)
    if not path.is_file():
        raise FileNotFoundError(f"Excel file not found: {filename}")

    wb = load_workbook(filename, read_only=True)
    try:
        ws = wb[sheet_name]
    except KeyError:
        wb.close()
        raise KeyError(f"Sheet '{sheet_name}' not found in {filename}") from None

    # Read header row to find column indices by name
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    col_map = {}
    for idx, value in enumerate(header_row):
        if value in REQUIRED_COLUMNS:
            col_map[value] = idx

    missing = REQUIRED_COLUMNS - col_map.keys()
    if missing:
        wb.close()
        raise ValueError(f"Missing required columns in Excel: {missing}")

    logger.info(f"Found columns: {col_map}")

    name_idx = col_map[NAME_COLUMN]
    url_idx = col_map["Document URL"]
    path_idx = col_map["Document Path"]

    records = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        name = row[name_idx] if name_idx < len(row) else None
        url = row[url_idx] if url_idx < len(row) else None
        file_path = row[path_idx] if path_idx < len(row) else None

        if not url or not file_path:
            logger.warning(f"Row {row_num}: missing URL or file path - skipping")
            continue

        records.append(
            {
                "name": str(name or ""),
                "url": str(url),
                "file_path": str(file_path),
                "row_number": row_num,
            }
        )

    wb.close()
    logger.info(f"Read {len(records)} document records from {filename}")
    return records
