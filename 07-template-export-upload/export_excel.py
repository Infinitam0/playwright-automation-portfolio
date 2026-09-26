"""Export document records to Excel using openpyxl."""

import logging

from openpyxl import Workbook
from openpyxl.styles import Font

logger = logging.getLogger(__name__)


def export_to_excel(
    documents: list[dict],
    filename: str,
    sheet_name: str,
    headers: list[dict] | None = None,
) -> None:
    """Write document records to an Excel file.

    When *headers* is provided (list of {"col_id", "label"} dicts from the
    AG-Grid), each grid column gets its own Excel column, followed by a
    "Document URL" and a "Document Path" column (the "_url" and "_file_path" keys).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    bold = Font(bold=True)

    if headers:
        col_ids = [h["col_id"] for h in headers]
        excel_headers = [h["label"] or h["col_id"] for h in headers] + ["Document URL", "Document Path"]

        for col, header_text in enumerate(excel_headers, 1):
            cell = ws.cell(row=1, column=col, value=header_text)
            cell.font = bold

        for row_idx, doc in enumerate(documents, 2):
            for col, col_id in enumerate(col_ids, 1):
                ws.cell(row=row_idx, column=col, value=doc.get(col_id, ""))
            ws.cell(row=row_idx, column=len(col_ids) + 1, value=doc.get("_url", ""))
            ws.cell(row=row_idx, column=len(col_ids) + 2, value=doc.get("_file_path", ""))
    else:
        # Fallback: legacy 2-column format
        for col, header_text in enumerate(["Document name", "Document URL"], 1):
            cell = ws.cell(row=1, column=col, value=header_text)
            cell.font = bold
        for row_idx, doc in enumerate(documents, 2):
            ws.cell(row=row_idx, column=1, value=doc.get("name", ""))
            ws.cell(row=row_idx, column=2, value=doc.get("url", ""))

    # Auto-width columns
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 100)

    wb.save(filename)
    logger.info(f"Exported {len(documents)} documents to {filename}")
