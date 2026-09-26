"""
export_excel.py

Incremental Excel writer for the calendar export.
Supports crash-resilient, resumable writes using openpyxl.
"""

import logging
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

logger = logging.getLogger(__name__)

COLUMNS = [
    "Appointment_ID",
    "Date",
    "Time",
    "Attendee name",
    "Time category",
    "Time_CategoryA",
    "Time_CategoryB",
    "Time_CategoryC",
    "Time_CategoryD",
    "Created by",
    "Created on",
]

# Reasonable default column widths (in characters) keyed by column name.
COLUMN_WIDTHS = {
    "Appointment_ID": 16,
    "Date": 14,
    "Time": 14,
    "Attendee name": 30,
    "Time category": 20,
    "Time_CategoryA": 20,
    "Time_CategoryB": 22,
    "Time_CategoryC": 22,
    "Time_CategoryD": 24,
    "Created by": 22,
    "Created on": 20,
}


class ExcelWriter:
    """Incremental Excel writer with resumability support."""

    def __init__(self, filepath: str, sheet_name: str) -> None:
        """
        If file exists: load it, build the set of already-exported
        Appointment_ID|attendee keys, and find the max Date for resume optimisation.
        If file does not exist: create a new workbook with bold headers.
        """
        self._filepath = filepath
        self._sheet_name = sheet_name
        self._already_exported: set[str] = set()
        self._max_exported_date: str | None = None

        if Path(filepath).exists():
            self._load_existing()
        else:
            self._create_new()

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    def _load_existing(self) -> None:
        """Load an existing workbook and populate the resumability state."""
        logger.info("Loading existing Excel file: %s", self._filepath)
        self._wb = load_workbook(self._filepath)

        if self._sheet_name in self._wb.sheetnames:
            self._ws = self._wb[self._sheet_name]
        else:
            logger.warning(
                "Sheet '%s' not found in %s; creating it.",
                self._sheet_name,
                self._filepath,
            )
            self._ws = self._wb.create_sheet(self._sheet_name)
            self._write_headers()

        # Iterate data rows (row 1 is the header).
        # Use composite key (Appointment_ID + attendee name) for dedup so that
        # multi-attendee appointments produce one row per attendee.
        for row in self._ws.iter_rows(min_row=2, values_only=True):
            appointment_id = str(row[0]) if row[0] is not None else ""
            date = str(row[1]) if row[1] is not None else ""
            attendee_name = str(row[3]) if row[3] is not None else ""
            if appointment_id:
                self._already_exported.add(f"{appointment_id}|{attendee_name}")
            if date and (self._max_exported_date is None or date > self._max_exported_date):
                self._max_exported_date = date

        logger.info(
            "Loaded %d existing rows; max date = %s",
            len(self._already_exported),
            self._max_exported_date,
        )

    def _create_new(self) -> None:
        """Create a fresh workbook with bold column headers."""
        logger.info("Creating new Excel file: %s", self._filepath)
        self._wb = Workbook()
        # openpyxl creates a default sheet named 'Sheet'; rename or replace it.
        default_sheet = self._wb.active
        default_sheet.title = self._sheet_name
        self._ws = default_sheet
        self._write_headers()
        self._wb.save(self._filepath)

    def _write_headers(self) -> None:
        """Write bold column headers and apply default column widths."""
        bold_font = Font(bold=True)
        for col_index, col_name in enumerate(COLUMNS, start=1):
            cell = self._ws.cell(row=1, column=col_index, value=col_name)
            cell.font = bold_font

        self._apply_column_widths()

    def _apply_column_widths(self) -> None:
        """Set reasonable column widths based on the predefined mapping."""
        for col_index, col_name in enumerate(COLUMNS, start=1):
            col_letter = self._ws.cell(row=1, column=col_index).column_letter
            width = COLUMN_WIDTHS.get(col_name, 15)
            self._ws.column_dimensions[col_letter].width = width

    # ------------------------------------------------------------------
    # Public properties
    # ------------------------------------------------------------------

    @property
    def already_exported(self) -> set[str]:
        """Return the set of Appointment_ID|attendee keys already in the Excel file."""
        return self._already_exported

    @property
    def max_exported_date(self) -> str | None:
        """Return the maximum Date string (YYYY-MM-DD) found in existing data, or None."""
        return self._max_exported_date

    @property
    def row_count(self) -> int:
        """Number of data rows, excluding the header row."""
        # max_row includes the header, so subtract 1. Returns 0 for an empty sheet.
        max_row = self._ws.max_row
        return max(0, max_row - 1)

    # ------------------------------------------------------------------
    # Public methods
    # ------------------------------------------------------------------

    def is_exported(self, appointment_id: str, attendee_name: str = "") -> bool:
        """Return True if this Appointment_ID + attendee combination is already exported."""
        return f"{appointment_id}|{attendee_name}" in self._already_exported

    def append_row(self, row_data: dict) -> None:
        """
        Append a single row to the worksheet and save the file immediately.

        row_data keys must match the column names defined in COLUMNS.
        After writing, the Appointment_ID|attendee key is added to already_exported
        and max_exported_date is updated if applicable.
        """
        row_values = [row_data.get(col) for col in COLUMNS]
        self._ws.append(row_values)

        appointment_id = str(row_data.get("Appointment_ID", "") or "")
        date = str(row_data.get("Date", "") or "")
        attendee_name = str(row_data.get("Attendee name", "") or "")
        if appointment_id:
            self._already_exported.add(f"{appointment_id}|{attendee_name}")

        if date and (self._max_exported_date is None or date > self._max_exported_date):
            self._max_exported_date = date

        self._wb.save(self._filepath)
        logger.debug(
            "Appended row (ID=%s, Date=%s); total rows=%d",
            appointment_id,
            date,
            self.row_count,
        )
