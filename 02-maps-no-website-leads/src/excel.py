"""Write results to a multi-sheet .xlsx.

Sheets are ordered by how worth calling the businesses are, so the file opens on
the ones to work first. "All Results" exists so a lead can be checked against
what was actually scraped rather than taken on faith.
"""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import PlaceDetail, WebsiteTier

logger = logging.getLogger(__name__)

COLUMNS = [
    ("name", "Business", 34),
    ("category", "Category", 22),
    ("phone", "Phone", 18),
    ("address", "Address", 44),
    ("website", "Website", 34),
    ("website_tier", "Tier", 17),
    ("rating", "Rating", 8),
    ("review_count", "Reviews", 9),
    ("maps_url", "Maps Link", 46),
    ("plus_code", "Plus Code", 24),
    ("latitude", "Lat", 11),
    ("longitude", "Lng", 11),
    ("cid", "CID", 22),
    ("place_id", "Place ID", 30),
    ("extraction_misses", "Extraction Misses", 20),
]

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

# Phone numbers and CIDs must stay text. Excel turns "+1 555 010 0199" into a
# formula-ish mess and silently truncates a 20-digit CID to float precision,
# which corrupts the one field that identifies the business.
_TEXT_COLUMNS = {"phone", "cid", "place_id", "plus_code"}


def _write_sheet(ws: Worksheet, rows: list[PlaceDetail]) -> None:
    ws.append([label for _, label, _ in COLUMNS])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for detail in rows:
        d = detail.to_row()
        ws.append([d.get(key) for key, _, _ in COLUMNS])

    for idx, (key, _, width) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = width
        if key in _TEXT_COLUMNS:
            for cell in ws[letter][1:]:
                cell.number_format = "@"

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"


def write_workbook(details: list[PlaceDetail], out_path: Path) -> dict[str, int]:
    by_tier: dict[WebsiteTier, list[PlaceDetail]] = {t: [] for t in WebsiteTier}
    for d in details:
        by_tier[d.website_tier].append(d)

    sheets = [
        ("No Website", by_tier[WebsiteTier.NONE]),
        ("Social Only", by_tier[WebsiteTier.SOCIAL_ONLY]),
        ("Ordering Platform", by_tier[WebsiteTier.ORDERING_PLATFORM]),
        ("All Results", details),
    ]

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        # Sort by review count desc: an established business with 200 reviews and
        # no website is a far better call than a listing with none.
        rows = sorted(rows, key=lambda d: (d.review_count or 0), reverse=True)
        _write_sheet(wb.create_sheet(title), rows)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    counts = {title: len(rows) for title, rows in sheets}
    logger.info("wrote %s  %s", out_path, counts)
    return counts
