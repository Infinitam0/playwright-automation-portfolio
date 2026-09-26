import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.excel import write_workbook  # noqa: E402
from src.models import Place, PlaceDetail, WebsiteTier  # noqa: E402
from src.store import Store  # noqa: E402


def _detail(fid, name, tier, website=None, reviews=0):
    return PlaceDetail(
        fid=fid, maps_url=f"https://maps.google.com/{fid}", name=name,
        website=website, website_tier=tier, review_count=reviews,
        phone="+1 555 010 0199", cid="12345678901234567890",
    )


@pytest.fixture
def store(tmp_path):
    with Store(tmp_path / "t.db") as s:
        yield s


def test_places_are_deduped_on_insert(store):
    p = Place(fid="0xa:0xb", maps_url="u1", name="X")
    assert store.add_places([p]) == 1
    assert store.add_places([p]) == 0
    assert len(store.pending_places()) == 1


def test_saving_a_detail_clears_it_from_pending(store):
    store.add_places([Place(fid="0xa:0xb", maps_url="u1", name="X")])
    assert len(store.pending_places()) == 1

    store.save_detail(_detail("0xa:0xb", "X", WebsiteTier.NONE))
    assert store.pending_places() == []
    assert store.counts() == {"done": 1}


def test_resume_skips_already_detailed_places(store):
    """The point of the database: an interrupted run must not re-fetch."""
    places = [Place(fid=f"0xa:0x{i}", maps_url=f"u{i}") for i in range(5)]
    store.add_places(places)

    for p in places[:3]:
        store.save_detail(_detail(p.fid, "done", WebsiteTier.REAL, "https://x.example"))

    remaining = [p.fid for p in store.pending_places()]
    assert remaining == ["0xa:0x3", "0xa:0x4"]


def test_failed_is_distinct_from_done(store):
    """A fetch failure says nothing about the business's website, so it must
    never look like a completed no-website result."""
    store.add_places([Place(fid="0xa:0xb", maps_url="u1")])
    store.mark_failed("0xa:0xb")

    assert store.counts() == {"failed": 1}
    assert store.pending_places() == []
    assert store.all_details() == []


def test_retry_failed_requeues(store):
    store.add_places([Place(fid="0xa:0xb", maps_url="u1")])
    store.mark_failed("0xa:0xb")
    assert store.reset_failed() == 1
    assert len(store.pending_places()) == 1


def test_tiles_are_not_reharvested(store):
    assert not store.tile_done("52.09,5.12@15z", "barbershop")
    store.record_tile("52.09,5.12@15z", "barbershop", 42, False)
    assert store.tile_done("52.09,5.12@15z", "barbershop")


def test_tile_cache_is_per_query(store):
    """A multi-category sweep must search each position once per query. Keying
    the cache on position alone would skip every tile after the first category
    and silently return almost nothing."""
    store.record_tile("52.09,5.12@15z", "hair salon", 42, False)
    assert store.tile_done("52.09,5.12@15z", "hair salon")
    assert not store.tile_done("52.09,5.12@15z", "nail salon")


def test_details_round_trip_through_the_database(store):
    store.add_places([Place(fid="0xa:0xb", maps_url="u1")])
    store.save_detail(
        _detail("0xa:0xb", "Nails", WebsiteTier.SOCIAL_ONLY, "https://instagram.com/x")
    )
    got = store.all_details()
    assert len(got) == 1
    assert got[0].website_tier is WebsiteTier.SOCIAL_ONLY
    assert got[0].name == "Nails"


def test_workbook_sheets_and_routing(tmp_path):
    details = [
        _detail("0x1:0x1", "No Site A", WebsiteTier.NONE, reviews=10),
        _detail("0x1:0x2", "No Site B", WebsiteTier.NONE, reviews=99),
        _detail("0x1:0x3", "Insta", WebsiteTier.SOCIAL_ONLY, "https://instagram.com/x"),
        _detail("0x1:0x4", "Booking", WebsiteTier.ORDERING_PLATFORM, "https://booksy.com/x"),
        _detail("0x1:0x5", "Real", WebsiteTier.REAL, "https://real.example"),
    ]
    out = tmp_path / "leads.xlsx"
    counts = write_workbook(details, out)

    assert counts == {
        "No Website": 2, "Social Only": 1,
        "Ordering Platform": 1, "All Results": 5,
    }

    wb = load_workbook(out)
    assert wb.sheetnames == ["No Website", "Social Only", "Ordering Platform", "All Results"]
    ws = wb["No Website"]
    assert ws["A1"].value == "Business"
    # Sorted by review count desc, so the best lead is the first row.
    assert ws["A2"].value == "No Site B"
    assert ws.freeze_panes == "A2"


def test_phone_and_cid_stay_text(tmp_path):
    """Excel turns a leading + into a formula and truncates a 20-digit CID to
    float precision, which silently corrupts the business identifier."""
    out = tmp_path / "leads.xlsx"
    write_workbook([_detail("0x1:0x1", "X", WebsiteTier.NONE)], out)

    ws = load_workbook(out)["No Website"]
    headers = [c.value for c in ws[1]]
    phone_col = headers.index("Phone") + 1
    cid_col = headers.index("CID") + 1

    assert ws.cell(row=2, column=phone_col).number_format == "@"
    assert ws.cell(row=2, column=cid_col).number_format == "@"
    assert ws.cell(row=2, column=cid_col).value == "12345678901234567890"


def test_empty_result_set_still_writes_a_valid_file(tmp_path):
    out = tmp_path / "empty.xlsx"
    counts = write_workbook([], out)
    assert counts["All Results"] == 0
    wb = load_workbook(out)
    assert wb["No Website"]["A1"].value == "Business"


def test_centre_is_pinned_across_resumes(store):
    """Geocoding is not deterministic; a drifting centre relabels every tile and
    makes a resumed run re-harvest the entire grid."""
    assert store.remembered_centre("Example City") is None
    store.remember_centre("Example City", 52.09070, 5.12140)
    assert store.remembered_centre("Example City") == (52.09070, 5.12140)
    assert store.remembered_centre("example city  ") == (52.09070, 5.12140)
    assert store.remembered_centre("Other Town") is None
