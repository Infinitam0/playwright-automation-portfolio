import logging
from pathlib import Path

from openpyxl import load_workbook

logger = logging.getLogger(__name__)


def read_person_names(file_path: str) -> list[str]:
    """Read person names from column B of the Excel file, skipping the header row."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Input Excel file not found: {path.resolve()}")

    wb = load_workbook(path, read_only=True)
    ws = wb.active

    names = []
    for row_idx, row in enumerate(ws.iter_rows(min_col=2, max_col=2, values_only=True), start=1):
        if row_idx == 1:
            continue  # skip header
        value = row[0]
        if value and str(value).strip():
            names.append(str(value).strip())

    wb.close()
    logger.info(f"Read {len(names)} person names from {path.name}")
    return names
