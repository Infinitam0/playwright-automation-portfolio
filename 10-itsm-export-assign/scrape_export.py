import logging
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from playwright.async_api import Frame

from models import AppConfig, EventType, PersonEvent

logger = logging.getLogger(__name__)

# Selectors are placeholders; adapt them to the target portal.
SELECTORS = {
    "person_event": ".history-item.person-unassigned, .history-item.person-assigned",
    "person_link": ".subject a",  # the affected person (detail link)
    "performer": ".actor",  # who performed the action
    "load_more_label": "Load more",  # UI label on the target site
}

# Month names for timestamp parsing (the event list writes dates out in words)
MONTHS = {
    "january": 1,
    "february": 2,
    "march": 3,
    "april": 4,
    "may": 5,
    "june": 6,
    "july": 7,
    "august": 8,
    "september": 9,
    "october": 10,
    "november": 11,
    "december": 12,
}


def parse_event_timestamp(text: str) -> datetime | None:
    """Parse datetime strings like '1 January 2000 00:00' or '01-01-2000 00:00'."""
    # Try DD-MM-YYYY HH:MM
    m = re.search(r"(\d{1,2})-(\d{1,2})-(\d{4})\s+(\d{1,2}):(\d{2})", text)
    if m:
        try:
            return datetime(
                year=int(m.group(3)),
                month=int(m.group(2)),
                day=int(m.group(1)),
                hour=int(m.group(4)),
                minute=int(m.group(5)),
            )
        except ValueError:
            pass

    # Try "D month YYYY HH:MM"
    pattern = r"(\d{1,2})\s+(" + "|".join(MONTHS.keys()) + r")\s+(\d{4})\s+(\d{1,2}):(\d{2})"
    m = re.search(pattern, text, re.IGNORECASE)
    if m:
        try:
            return datetime(
                year=int(m.group(3)),
                month=MONTHS[m.group(2).lower()],
                day=int(m.group(1)),
                hour=int(m.group(4)),
                minute=int(m.group(5)),
            )
        except ValueError:
            pass

    return None


def classify_event(raw_text: str) -> EventType | None:
    """Return the link/unlink type of an event's text, or None for other events."""
    raw_lower = raw_text.lower()
    # Check "unassigned" first: it contains "assigned".
    if "unassigned" in raw_lower:
        return EventType.UNASSIGNED
    if "assigned" in raw_lower:
        return EventType.ASSIGNED
    return None


async def load_all_events(events_frame: Frame, page, config: AppConfig) -> None:
    """Click 'Load more' until all events are loaded."""
    logger.info("Loading all events...")
    clicks = 0
    while clicks < config.portal.max_load_more_clicks:
        try:
            btn = events_frame.get_by_role("button", name=SELECTORS["load_more_label"])
            await btn.wait_for(state="visible", timeout=3000)
            await btn.click()
            clicks += 1
            logger.info(f"'Load more' click #{clicks}")
            await page.wait_for_timeout(config.portal.load_more_wait_ms)
        except Exception:
            logger.info(f"'Load more' button gone after {clicks} clicks - all events loaded.")
            break
    logger.info(f"Total 'Load more' clicks: {clicks}")


async def scrape_events(events_frame: Frame, config: AppConfig) -> list[PersonEvent]:
    """Scrape person events from the events frame and filter."""
    logger.info("Scraping events...")
    # Select both assigned and unassigned person events
    event_elements = events_frame.locator(SELECTORS["person_event"])
    count = await event_elements.count()
    logger.info(f"Found {count} person event elements.")

    events: list[PersonEvent] = []
    for i in range(count):
        el = event_elements.nth(i)
        raw_text = await el.inner_text()

        # Determine event type from the description text
        event_type = classify_event(raw_text)
        if event_type is None:
            continue  # not an assign/unassign event we care about

        # Person name is in the detail link, the performer in its own element
        person_name = ""
        performed_by = ""
        try:
            person_link = el.locator(SELECTORS["person_link"])
            if await person_link.count() > 0:
                person_name = (await person_link.first.inner_text()).strip()
        except Exception:
            pass
        try:
            user_el = el.locator(SELECTORS["performer"])
            if await user_el.count() > 0:
                performed_by = (await user_el.first.inner_text()).strip()
        except Exception:
            pass

        # Parse timestamp
        timestamp = parse_event_timestamp(raw_text)

        events.append(
            PersonEvent(
                event_type=event_type,
                person_name=person_name,
                timestamp=timestamp,
                performed_by=performed_by,
                raw_text=raw_text.strip().replace("\n", " | "),
            )
        )

    logger.info(f"Scraped {len(events)} person events total.")

    # Filter
    filtered = []
    for ev in events:
        if ev.event_type != config.filter.event_type:
            continue
        if config.filter.filter_date:
            fd = config.filter.filter_date
            if ev.timestamp is None:
                continue
            if ev.timestamp.replace(second=0, microsecond=0) != fd.replace(second=0, microsecond=0):
                continue
        filtered.append(ev)

    logger.info(
        f"Filtered to {len(filtered)} '{config.filter.event_type.value}' events"
        + (f" at {config.filter.filter_date}" if config.filter.filter_date else "")
    )
    return filtered


def export_to_excel(events: list[PersonEvent], output_file: str) -> None:
    """Write events to Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Events"

    # Column B ("Person") is what reassign_people/read_excel.py reads back
    headers = ["Type", "Person", "Date/Time", "Performed by", "Raw text"]
    bold = Font(bold=True)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = bold

    for row_idx, ev in enumerate(events, 2):
        ws.cell(row=row_idx, column=1, value=ev.event_type.value)
        ws.cell(row=row_idx, column=2, value=ev.person_name)
        ws.cell(row=row_idx, column=3, value=ev.timestamp.strftime("%d-%m-%Y %H:%M") if ev.timestamp else "")
        ws.cell(row=row_idx, column=4, value=ev.performed_by)
        ws.cell(row=row_idx, column=5, value=ev.raw_text)

    # Auto-width columns
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 80)

    wb.save(output_file)
    logger.info(f"Exported {len(events)} events to {output_file}")
