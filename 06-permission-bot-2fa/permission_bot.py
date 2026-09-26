"""
Permission bot: logs in to a SaaS permissions portal (password +
TOTP 2FA) and checks/unchecks access permissions per account, driven by
permissions.xlsx. Every run appends one SUCCESS/FAILURE line to permission_bot.log.
"""
import logging
import os
import sys
from datetime import datetime
from itertools import groupby
from operator import itemgetter
from pathlib import Path

import pyotp
from dotenv import load_dotenv
from openpyxl import load_workbook
from playwright.sync_api import Locator, Page, sync_playwright

import labels as ui

# =============================================================================
# Configuration from environment variables (loaded from .env if present)
# =============================================================================

# Load .env file from same directory as this script
_env_path = Path(__file__).parent / ".env"
load_dotenv(_env_path)

# Script directory (used for log file path)
_script_dir = Path(__file__).parent.resolve()


def append_log_line(status: str, detail: str, log_dir: Path = None):
    """
    Append exactly one line to the persistent log file.

    Format: 2025-01-01 12:34:56 | SUCCESS | 3 accounts, 12 actions, 0 failures
    """
    target_dir = log_dir or _script_dir
    target_dir.mkdir(parents=True, exist_ok=True)
    log_path = target_dir / "permission_bot.log"
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line = f"{timestamp} | {status} | {detail}\n"
    with open(log_path, "a", encoding="utf-8") as f:
        f.write(line)


def get_config():
    """
    Load configuration from environment variables.

    Required:
        PORTAL_USER: Login username
        PORTAL_PASS: Login password
        PORTAL_TOTP_SECRET: Base32 TOTP secret for 2FA code generation

    Optional:
        HEADLESS: "true" (default) or "false" for headed mode
        LOG_DIR: Directory for log files (default: script directory)
        LOGIN_URL: Login URL (default: $PORTAL_BASE_URL/admin)

    Returns:
        dict with configuration values

    Raises:
        SystemExit if required variables are missing
    """
    missing = []

    username = os.environ.get("PORTAL_USER")
    if not username:
        missing.append("PORTAL_USER")

    password = os.environ.get("PORTAL_PASS")
    if not password:
        missing.append("PORTAL_PASS")

    totp_secret = os.environ.get("PORTAL_TOTP_SECRET", "").strip().replace(" ", "")
    if not totp_secret:
        missing.append("PORTAL_TOTP_SECRET")

    if missing:
        detail = f"Missing env vars: {', '.join(missing)}"
        print(f"ERROR: {detail}", file=sys.stderr)
        append_log_line("FAILURE", detail)
        sys.exit(1)

    # Parse headless setting (default: True for production/scheduled runs)
    headless_str = os.environ.get("HEADLESS", "true").lower()
    headless = headless_str not in ("false", "0", "no")

    # Log directory (default: script directory)
    log_dir = os.environ.get("LOG_DIR")
    if log_dir:
        log_dir = Path(log_dir).resolve()
    else:
        log_dir = _script_dir

    # Ensure log directory exists
    log_dir.mkdir(parents=True, exist_ok=True)

    # Login URL
    base_url = (os.environ.get("PORTAL_BASE_URL") or "https://portal.example.com").rstrip("/")
    url = os.environ.get("LOGIN_URL") or f"{base_url}/admin"

    return {
        "username": username,
        "password": password,
        "totp_secret": totp_secret,
        "headless": headless,
        "log_dir": log_dir,
        "url": url,
    }


# Load configuration early (before logging setup)
CONFIG = get_config()

# Console-only logging (no file handler -- persistent log uses append_log_line)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    handlers=[logging.StreamHandler()],
)
logger = logging.getLogger(__name__)


def safe_check(locator, name: str):
    """Check a checkbox only if it's currently unchecked."""
    if not locator.is_checked():
        locator.check()
    else:
        logger.debug(f"Already checked: {name}")


def safe_uncheck(locator, name: str):
    """Uncheck a checkbox only if it's currently checked."""
    if locator.is_checked():
        locator.uncheck()
    else:
        logger.debug(f"Already unchecked: {name}")


def validate_permissions_file(filepath: str) -> list[dict]:
    """
    Read and validate permission configurations from Excel file.

    Args:
        filepath: Path to the Excel file containing permission data

    Returns:
        List of validated permission dictionaries

    Raises:
        FileNotFoundError: If the Excel file doesn't exist
        ValueError: If headers are invalid or no valid rows found
    """
    path = Path(filepath)

    if not path.exists():
        raise FileNotFoundError(f"Permission file not found: {filepath}")

    wb = load_workbook(filename=path, read_only=True, data_only=True)
    ws = wb.active

    expected_headers = ["AccountName", "ParentPath", "CheckboxName", "Action"]
    actual_headers = [cell.value for cell in ws[1]]
    if actual_headers != expected_headers:
        wb.close()
        raise ValueError(f"Invalid headers. Expected {expected_headers}, got {actual_headers}")

    permissions = []

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        account, parent_path, checkbox, action = row

        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        if account is None or str(account).strip() == "":
            logger.warning(f"Row {row_num}: Missing AccountName, skipping")
            continue

        if checkbox is None or str(checkbox).strip() == "":
            logger.warning(f"Row {row_num}: Missing CheckboxName, skipping")
            continue

        if action is None or str(action).strip() == "":
            logger.warning(f"Row {row_num}: Missing Action, skipping")
            continue

        action_lower = str(action).strip().lower()
        if action_lower not in ["check", "uncheck"]:
            logger.warning(f"Row {row_num}: Invalid Action '{action}', expected 'check' or 'uncheck', skipping")
            continue

        path_list = None
        if parent_path is not None and str(parent_path).strip() != "":
            path_list = [node.strip() for node in str(parent_path).split(" > ")]

        permissions.append({
            "account": str(account).strip(),
            "parent_path": path_list,
            "checkbox": str(checkbox).strip(),
            "action": action_lower,
            "row_num": row_num,
        })

    wb.close()

    if not permissions:
        raise ValueError("No valid permission rows found in Excel file")

    logger.info(f"Loaded {len(permissions)} permission entries from Excel")
    return permissions


def group_permissions_by_account(permissions: list[dict]) -> dict[str, list[dict]]:
    """Group permissions by account name for batch processing."""
    permissions.sort(key=itemgetter("account"))
    return {
        account: list(group)
        for account, group in groupby(permissions, key=itemgetter("account"))
    }


def process_account(page: Page, account_name: str, actions: list[dict]) -> tuple[int, list[dict]]:
    """
    Process all permission actions for a single account using path-based navigation.

    Args:
        page: Playwright page object (the admin page)
        account_name: User to select in tree
        actions: List of permission dicts with parent_path (list) for this account

    Returns:
        Tuple of (successful_count, failed_actions_list)
    """
    successful_count = 0
    failed_actions = []

    # Select user in tree
    select_user_in_tree(page, account_name)
    page.wait_for_timeout(500)

    # Open the permissions tab
    page.get_by_text(ui.TAB_PERMISSIONS, exact=True).click()
    page.wait_for_timeout(500)

    # Open the access editor
    page.get_by_role("button", name=ui.BTN_MANAGE_ACCESS).locator("a").click()
    page.wait_for_load_state("networkidle")

    # Wait for tree to load
    page.locator("[role='treeitem']").first.wait_for(state="visible", timeout=10000)
    page.wait_for_timeout(500)

    # Process each action using path-based navigation
    for action_dict in actions:
        parent_path = action_dict["parent_path"]
        checkbox_name = action_dict["checkbox"]

        try:
            if parent_path:
                current_context = page

                for node_name in parent_path:
                    node = current_context.get_by_role("treeitem", name=node_name).first

                    if node.count() == 0:
                        raise Exception(f"Path node not found: '{node_name}'")

                    expanded_attr = node.get_attribute("aria-expanded")
                    if expanded_attr == "false":
                        expand_icon = node.locator("i").first
                        if expand_icon.is_visible():
                            expand_icon.click()
                            first_child = node.locator("[role='treeitem']").first
                            first_child.wait_for(state="visible", timeout=5000)
                        else:
                            raise Exception(f"Cannot expand '{node_name}': icon not visible")

                    current_context = node

                checkbox_treeitem = current_context.get_by_role("treeitem", name=checkbox_name).first
            else:
                checkbox_treeitem = page.get_by_role("treeitem", name=checkbox_name).first

            if checkbox_treeitem.count() == 0:
                raise Exception(f"Checkbox treeitem not found: '{checkbox_name}'")

            checkbox_locator = checkbox_treeitem.locator("input[type='checkbox']")
            if checkbox_locator.count() == 0:
                checkbox_locator = checkbox_treeitem.get_by_role("checkbox")
            if checkbox_locator.count() == 0:
                checkbox_locator = checkbox_treeitem

            if action_dict["action"] == "check":
                safe_check(checkbox_locator, checkbox_name)
            else:
                safe_uncheck(checkbox_locator, checkbox_name)

            successful_count += 1

        except Exception as e:
            logger.error(f"  Row {action_dict['row_num']}: {action_dict['action']} '{checkbox_name}' -- FAILED: {e}")
            failed_actions.append({
                "row_num": action_dict["row_num"],
                "checkbox": checkbox_name,
                "error": str(e),
            })

    # Save changes, then confirm (nothing to save when every action failed)
    if successful_count:
        page.get_by_role("button", name=ui.BTN_SAVE).locator("a").click()
        page.wait_for_timeout(1000)
        page.get_by_role("button", name=ui.BTN_OK).locator("a").click()
        page.wait_for_timeout(500)
    else:
        close_editor(page)

    return (successful_count, failed_actions)


def close_editor(page: Page) -> None:
    """Close the access editor without saving (Close button, else the dialog's X)."""
    close_btn = page.get_by_role("button", name=ui.BTN_CLOSE)
    if close_btn.count() > 0 and close_btn.first.is_visible():
        close_btn.first.click()
    else:
        page.locator(ui.DIALOG_CLOSE_X).first.click()
    page.wait_for_timeout(500)


def process_all_accounts(page: Page, grouped: dict[str, list[dict]]) -> tuple[int, int, list[dict]]:
    """
    Process all accounts with progress tracking and error aggregation.

    Returns:
        Tuple of (total_actions_succeeded, total_actions_count, all_failures_list)
    """
    failures = []
    total_accounts = len(grouped)
    total_actions = sum(len(actions) for actions in grouped.values())
    total_succeeded = 0

    for idx, (account_name, actions) in enumerate(grouped.items(), start=1):
        try:
            succeeded, action_failures = process_account(page, account_name, actions)
            total_succeeded += succeeded

            if action_failures:
                failures.append({
                    "account": account_name,
                    "type": "action_failures",
                    "actions": action_failures,
                })

            logger.info(f"Account {idx}/{total_accounts}: {account_name} -- {succeeded}/{len(actions)} succeeded")

        except Exception as e:
            failures.append({
                "account": account_name,
                "type": "account_error",
                "error": str(e),
                "skipped_actions": len(actions),
            })
            logger.error(f"Account {idx}/{total_accounts}: {account_name} -- FAILED: {e}")
            continue

    return (total_succeeded, total_actions, failures)


def select_user_in_tree(page: Page, user_name: str):
    """Select a user in the tree view, scrolling if necessary."""
    title_with_space = ui.TREE_NODE_TITLE.format(name=user_name)
    user_locator = page.locator(f'{ui.TREE_NODE}[title="{title_with_space}"]')
    user_locator.wait_for(state="attached", timeout=10000)
    user_locator.scroll_into_view_if_needed()
    page.wait_for_timeout(200)
    user_locator.click()


def is_node_expanded(node: Locator) -> bool:
    """Check if a tree node is currently expanded."""
    expanded_attr = node.get_attribute("aria-expanded")
    if expanded_attr is None:
        return False
    return expanded_attr == "true"


def expand_tree_node(page: Page, node_name: str, parent_name: str = None) -> None:
    """
    Expand a tree node by name, handling already-expanded state (idempotent).
    """
    title_with_space = ui.TREE_NODE_TITLE.format(name=node_name)
    user_node_div = page.locator(f'{ui.TREE_NODE}[title="{title_with_space}"]')

    if user_node_div.count() > 0:
        node = user_node_div.locator("xpath=ancestor::li[@role='treeitem'][1]")
    else:
        node = page.get_by_role("treeitem", name=node_name, exact=False).first

    node.wait_for(state="attached", timeout=10000)

    collapsed_ancestors = node.locator("xpath=ancestor::li[@role='treeitem'][@aria-expanded='false']")
    ancestor_count = collapsed_ancestors.count()

    if ancestor_count > 0:
        for i in range(ancestor_count):
            ancestor = collapsed_ancestors.nth(i)
            ancestor.locator("i").first.click()
            page.wait_for_timeout(200)

    if is_node_expanded(node):
        return

    node.locator("i").first.click()

    first_child = node.locator("[role='treeitem']").first
    first_child.wait_for(state="visible", timeout=5000)


def setup_browser_and_navigate(page: Page) -> Page:
    """
    Log in directly to the admin area, open user management, and expand the user tree.

    Args:
        page: Playwright page object

    Returns:
        Same page object after navigation setup
    """
    page.goto(CONFIG["url"])
    page.wait_for_load_state("networkidle")

    # Login if redirected to login page
    username_field = page.locator(ui.LOGIN_USERNAME)
    if username_field.is_visible():
        username_field.click()
        username_field.fill(CONFIG["username"])
        page.locator(ui.LOGIN_PASSWORD).click()
        page.locator(ui.LOGIN_PASSWORD).fill(CONFIG["password"])
        page.locator(ui.LOGIN_SUBMIT).click()
        page.wait_for_load_state("networkidle")

        # Handle 2FA
        code_input = page.locator(ui.LOGIN_OTP_CODE)
        try:
            code_input.wait_for(state="visible", timeout=3000)
            totp_code = pyotp.TOTP(CONFIG["totp_secret"]).now()
            code_input.click()
            code_input.fill(totp_code)
            page.locator(ui.LOGIN_SUBMIT).click()
        except Exception:
            pass  # No 2FA required

    # Wait for the admin page
    users_btn = page.get_by_role("button", name=ui.BTN_USERS).locator("a")
    users_btn.wait_for(state="visible", timeout=60000)

    # Navigate to user management
    users_btn.click()
    page.wait_for_load_state("networkidle")
    page.get_by_role("button", name=ui.BTN_USER_MANAGEMENT).locator("a").click()
    page.wait_for_load_state("networkidle")

    # Expand user tree (group names are specific to your portal tenant)
    expand_tree_node(page, "Org A Users")
    expand_tree_node(page, "Managers")

    return page


def main() -> int:
    """
    Main entry point for the permission bot.

    Returns:
        0 on success, 1 on failure
    """
    log_dir = CONFIG["log_dir"]

    # Phase 1: Validate permissions file
    try:
        script_dir = Path(__file__).parent
        permissions = validate_permissions_file(str(script_dir / "permissions.xlsx"))
    except (FileNotFoundError, ValueError) as e:
        logger.error(str(e))
        append_log_line("FAILURE", f"Permission file error: {e}", log_dir)
        return 1

    # Phase 2: Group by account
    grouped = group_permissions_by_account(permissions)
    total_accounts = len(grouped)
    total_actions = sum(len(actions) for actions in grouped.values())

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=CONFIG["headless"])
        context = browser.new_context()
        page = context.new_page()

        try:
            setup_browser_and_navigate(page)
            total_succeeded, total_actions, failures = process_all_accounts(page, grouped)

        except Exception as e:
            logger.error(f"Script error: {e}")
            append_log_line("FAILURE", f"Script error: {e}", log_dir)
            return 1
        finally:
            context.close()
            browser.close()

    # Final result
    failure_count = total_actions - total_succeeded
    if failures:
        detail = f"{total_accounts} accounts, {total_actions} actions, {failure_count} failure(s)"
        append_log_line("FAILURE", detail, log_dir)
        logger.warning(f"FAILURE | {detail}")
        return 1
    else:
        detail = f"{total_accounts} accounts, {total_actions} actions, 0 failures"
        append_log_line("SUCCESS", detail, log_dir)
        logger.info(f"SUCCESS | {detail}")
        return 0


if __name__ == "__main__":
    sys.exit(main())
